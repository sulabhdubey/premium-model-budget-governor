from io import BytesIO
import subprocess
from zipfile import ZipFile, ZIP_DEFLATED

import pytest

pytest.importorskip("docx")
pytest.importorskip("defusedxml")
from docx import Document

from premium_model_budget_governor.document_extract import extract_docx


def document_bytes(text="Review the contract and list unresolved questions."):
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Requirement"
    table.cell(0, 1).text = "Evidence"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def altered(data, name, content):
    result = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(result, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            if item.filename != name:
                target.writestr(item, source.read(item))
        target.writestr(name, content)
    return result.getvalue()


def test_real_worker_extracts_paragraphs_and_tables_with_provenance():
    result = extract_docx(document_bytes())
    assert "Review the contract" in result["text"]
    assert "Requirement\tEvidence" in result["text"]
    assert result["parser"].startswith("python-docx/")
    assert len(result["text_sha256"]) == 64
    assert result["truncated"] is False


@pytest.mark.parametrize("name,content", [
    ("../escape.xml", b"<xml/>"),
    ("word/vbaProject.bin", b"macro"),
    ("word/embeddings/object.bin", b"object"),
    ("word/header1.xml", b"<header/>"),
    ("word/document.xml", b'<!DOCTYPE x [<!ENTITY e "value">]><x>&e;</x>'),
    ("word/_rels/document.xml.rels", b'<Relationships><Relationship TargetMode="External" Target="https://example.invalid"/></Relationships>'),
])
def test_unsupported_document_content_is_withheld(name, content):
    with pytest.raises(ValueError):
        extract_docx(altered(document_bytes(), name, content))


def test_compression_bomb_and_secret_text_are_withheld():
    with pytest.raises(ValueError):
        extract_docx(altered(document_bytes(), "word/bomb.xml", b"x" * 3_000_000))
    with pytest.raises(ValueError):
        extract_docx(document_bytes("-----BEGIN PRIVATE KEY-----"))


def test_timeout_reports_no_partial_output(monkeypatch):
    import premium_model_budget_governor.document_extract as module
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("worker", 15)
    monkeypatch.setattr(module.subprocess, "run", timeout)
    with pytest.raises(ValueError):
        extract_docx(document_bytes())


def test_workbench_document_preview_and_source_change_gate(tmp_path):
    from premium_model_budget_governor.workbench import Workbench
    root = tmp_path / "project"
    root.mkdir()
    probe = lambda root: {"models": [{"model": "gpt-6-astra", "reasoning_efforts": ["low"], "input_modalities": ["text"]}]}
    app = Workbench({"sample": root}, tmp_path / "data", probe=probe)
    path = root / "brief.docx"
    path.write_bytes(document_bytes())
    preview = app.preview({"project": "sample", "mode": "astra_preferred", "budget_credits": 20,
                           "evidence": ["brief.docx"], "task": "Summarize this document with unanswered questions."})
    assert preview["plan"]["selected"]["stages"][0]["model"] == "gpt-6-astra"
    assert "Requirement" in preview["documents"][0]["text"]
    assert app.history() == []
    path.write_bytes(document_bytes("Changed source."))
    with pytest.raises(ValueError, match="changed"):
        app.execute(preview["id"], approved=True)
    assert app.history() == []


def test_worker_uses_private_copy_without_caller_environment_and_removes_it(monkeypatch):
    import premium_model_budget_governor.document_extract as module
    from pathlib import Path
    original = module.subprocess.run
    directories = []
    monkeypatch.setenv("PRIVATE_TEST_CREDENTIAL", "not-forwarded")
    def observed(*args, **kwargs):
        directory = Path(kwargs["cwd"])
        assert (directory / "input.docx").is_file()
        assert "PRIVATE_TEST_CREDENTIAL" not in kwargs["env"]
        directories.append(directory)
        return original(*args, **kwargs)
    monkeypatch.setattr(module.subprocess, "run", observed)
    extract_docx(document_bytes())
    assert directories and not directories[0].exists()


@pytest.mark.parametrize("markup", ["<w:sdt/>", "<w:ins/>", "<w:drawing/>", "<w:instrText>FIELD</w:instrText>"])
def test_unsupported_body_constructs_are_not_silently_dropped(markup):
    xml = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + markup + '</w:body></w:document>').encode()
    with pytest.raises(ValueError):
        extract_docx(altered(document_bytes(), "word/document.xml", xml))
