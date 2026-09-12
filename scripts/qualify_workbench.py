"""Run existing rendered Workbench QA on an isolated server that cannot dispatch."""
import argparse
import json
import os
from pathlib import Path
import subprocess
from threading import Thread

from premium_model_budget_governor.local_server import LocalServer
from premium_model_budget_governor.private_file import write_private, remove_if_unchanged
from premium_model_budget_governor.workbench import Workbench


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--node", required=True)
    parser.add_argument("--node-modules", required=True)
    parser.add_argument("--documents", action="store_true", help="Qualify optional DOCX intake using generated local fixtures")
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)

    def forbidden_executor(*args, **kwargs):
        raise RuntimeError("QA server cannot dispatch model calls")

    projects = {"examples": root / "examples"}
    if args.documents:
        from docx import Document
        folder = output / "document-fixtures"
        folder.mkdir()
        document = Document()
        document.add_paragraph("Draft project brief: review the requirements and list unresolved questions.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Requirement"
        table.cell(0, 1).text = "Evidence pending"
        document.save(folder / "brief.docx")
        (folder / "unsupported.pdf").write_bytes(b"%PDF-1.4\nUnsupported-format test fixture, not a valid PDF.")
        projects["documents"] = folder
    app = Workbench(projects, output / "data", executor=forbidden_executor)
    server = LocalServer(app)
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    session = output / "session.json"
    content = json.dumps({"url": server.origin + "/#session=" + server.token}).encode()
    write_private(session, content)
    env = {**os.environ, "NODE_PATH": args.node_modules, "PM_BG_SESSION_FILE": str(session),
           "PM_BG_DOCUMENT_QA": "1" if args.documents else "0",
           "PM_BG_QA_OUTPUT": str(output / "screenshots")}
    try:
        result = subprocess.run([args.node, str(root / "scripts/workbench_qa.cjs")],
                                cwd=root, env=env, timeout=600, check=False)
        return result.returncode
    finally:
        server.shutdown()
        server.server_close()
        thread.join(10)
        remove_if_unchanged(session, content)


if __name__ == "__main__":
    raise SystemExit(main())
